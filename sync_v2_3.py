#!/usr/bin/env python3
"""Sync final test report to all PPTs and MDs."""
import os, glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

BASE = r"D:\现代制造\Python编程实战"

PPT_CONTENT = [
    "JupyterHub v2.3 OAuth修复+测试报告",
    "",
    "[OAuth 400 修复]",
    "  根因: Hub重启导致OAuth客户端丢失",
    "  修复: 清除stale oauth + 重新登录",
    "  结果: teacher-zhang登录成功",
    "",
    "[数据保留]",
    "  PVC: claim-teacher-zhang Bound",
    "  76 Notebook + 8 Code + 2 Guides",
    "  结论: 零数据丢失",
    "",
    "[LLM测试]",
    "  master: 59s | teacher pod: 28s",
    "  模型: qwen2.5-coder:7b",
    "  瓶颈: CPU-only推理",
    "",
    "[CRDB测试]",
    "  3节点 Running | 27+数据库",
    "  读写分离: 26257/26267",
    "",
    "[资源优化]",
    "  milvus/nebula/pulsar stopped",
    "  Worker CPU: 79% -> 57%",
    "",
    "[账户体系]",
    "  33教师 | 43分组 | 多教师班级",
]

def update_ppts():
    ppts = [f for f in glob.glob(os.path.join(BASE, "**", "*.pptx"), recursive=True)
            if "_docx_work" not in f and "venv" not in f]
    updated = 0
    for path in ppts:
        try:
            prs = Presentation(path)
            has = False
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and "OAuth修复" in shape.text_frame.text:
                        has = True; break
                if has: break
            if has: continue
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
            tb.text_frame.text = PPT_CONTENT[0]
            p = tb.text_frame.paragraphs[0]
            p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F,0x49,0x7D)
            cb = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(6))
            ctf = cb.text_frame; ctf.word_wrap = True
            for i, line in enumerate(PPT_CONTENT[1:]):
                p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
                p.text = line
                if line.startswith("["):
                    p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F,0x49,0x7D)
                else:
                    p.font.size = Pt(10); p.font.color.rgb = RGBColor(0x33,0x33,0x33)
            prs.save(path)
            updated += 1
        except: pass
    return updated, len(ppts)

def update_mds():
    mds = []
    for root, dirs, files in os.walk(BASE):
        if "_docx_work" in root or "venv" in root or "node_modules" in root: continue
        for f in files:
            if f.endswith(".md"): mds.append(os.path.join(root, f))
    
    md_content = "\n\n---\n\n## 23. OAuth 400 修复 + 测试报告 (v2.3)\n\n### OAuth 400 修复\n- 根因: Hub重启导致OAuth客户端丢失\n- 修复: 清除stale oauth codes + 重新登录\n- 验证: teacher-zhang登录成功\n\n### 数据保留\n- PVC Bound | 76 Notebook | 8 Code | 双指南 | 评分系统\n- 结论: 零数据丢失\n\n### LLM测试\n- master: 59s | teacher pod: 28s | qwen2.5-coder:7b\n- 瓶颈: CPU-only推理\n\n### CRDB测试\n- 3节点 Running | 27+数据库 | 读写分离\n\n### 资源优化\n- milvus/nebula/pulsar stopped | Worker CPU 79% -> 57%\n"
    
    updated = 0
    for path in mds:
        try:
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
            if "OAuth 400" in content: continue
            with open(path, "a", encoding="utf-8") as f:
                f.write(md_content)
            updated += 1
        except: pass
    return updated, len(mds)

def update_guides():
    updated = 0
    for fname in ["JUPYTERHUB-OPERATION-GUIDE.md", "JUPYTERHUB-STUDENT-GUIDE.md"]:
        path = os.path.join(r"D:\dify-install", fname)
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
            if "OAuth 400" not in content:
                md_content = "\n\n---\n\n## 23. OAuth 400 修复 + 测试报告 (v2.3)\n\n### OAuth 400 修复\n- 根因: Hub重启导致OAuth客户端丢失\n- 修复: 清除stale oauth codes + 重新登录\n- 验证: teacher-zhang登录成功\n\n### 数据保留\n- PVC Bound | 76 Notebook | 8 Code | 双指南 | 评分系统\n- 结论: 零数据丢失\n\n### LLM测试\n- master: 59s | teacher pod: 28s | qwen2.5-coder:7b\n- 瓶颈: CPU-only推理\n\n### CRDB测试\n- 3节点 Running | 27+数据库 | 读写分离\n\n### 资源优化\n- milvus/nebula/pulsar stopped | Worker CPU 79% -> 57%\n"
                with open(path, "a", encoding="utf-8") as f:
                    f.write(md_content)
                updated += 1
    return updated

if __name__ == "__main__":
    print("=" * 70)
    print("Sync v2.3: OAuth fix + test report")
    print("=" * 70)
    
    pu, pt = update_ppts()
    print(f"PPT: {pu} updated / {pt} total")
    
    mu, mt = update_mds()
    print(f"MD: {mu} updated / {mt} total")
    
    gu = update_guides()
    print(f"Guides: {gu} updated")
    
    total = pu + mu + gu
    print(f"\nTotal: {total} files updated")
    print("=" * 70)
