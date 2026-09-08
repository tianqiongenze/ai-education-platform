#!/usr/bin/env python3
"""One-shot update: guides + PPTs + MDs with multi-language review, CRDB K8s, guides distribution."""
import os, glob, json

BASE = r"D:\现代制造\Python编程实战"
DIFY = r"D:\dify-install"

# Content to append to guides
GUIDE_APPEND = """
---

## 19. 多语言代码审查系统 (v2.1 新增)

### 支持语言与工具

| 语言 | 工具 | 检查内容 |
|------|------|----------|
| Python | pycodestyle + pyflakes | PEP8 规范、语法错误 |
| Java | Checkstyle / javac | 代码规范、编译错误 |
| Go | gofmt + go vet | 格式化、静态分析 |
| Rust | cargo clippy + rustfmt | 代码质量、格式化 |
| C/C++ | cppcheck | 内存泄漏、未初始化变量 |
| JavaScript | node --check | 语法错误 |

### 使用方式

```bash
# 单文件审查
python3 code_review.py src/main.py

# 目录审查（自动检测语言，递归所有源文件）
python3 code_review.py /home/jovyan/work/industrial-analytics/src/
python3 code_review.py /home/jovyan/work/mes-system/production-service/src/
python3 code_review.py /home/jovyan/work/industrial-gateway/
python3 code_review.py /home/jovyan/work/security-audit/
```

### 输出格式（JSON）

```json
{
  "language": "Python",
  "tool": "pycodestyle",
  "issues": ["src/main.py:10:1: E302 expected 2 blank lines"],
  "score": 85
}
```

### 评分规则

| 语言 | 每个问题扣分 | 满分 |
|------|-------------|------|
| Python | 100/总行数 | 100 |
| Java | 5分/问题 | 100 |
| Go | 10分/问题 | 100 |
| Rust | 5分/问题 | 100 |
| C/C++ | 5分/问题 | 100 |

---

## 20. 指南分发规则 (v2.1 新增)

| 用户类型 | 教师版指南 | 学生版指南 | 评分系统 | AI助手 |
|----------|-----------|-----------|----------|--------|
| teacher-zhang | ✅ | ✅ | ✅ | ✅ |
| Lecture-* (16个) | ✅ | ✅ | ✅ | ✅ |
| 学生（任何注册） | ❌ | ✅ | ✅ | ✅ |

---

## 21. CockroachDB K8s 部署确认 (v2.1 新增)

### 部署方式

CockroachDB 已从裸金属迁移到 K8s Pod 部署：

| 项目 | 旧（裸金属） | 新（K8s Pod） |
|------|-------------|--------------|
| 二进制路径 | /opt/cockroach/cockroach | /cockroach-binary/cockroach（容器内） |
| 数据路径 | /home/crdb-data/a | /cockroach/cockroach-data（hostPath） |
| systemd | cockroach-node1/2.service | StatefulSet |
| systemd 状态 | — | disabled + inactive (dead) |
| 管理方式 | systemctl | kubectl / Rancher |
| 网络模式 | 宿主机直接 | hostNetwork: true |

> **说明**：由于使用 `hostNetwork: true`，Pod 进程在宿主机 `ps aux` 中可见，
> 但实际运行在容器内（二进制路径为 `/cockroach-binary/`，非 `/opt/cockroach/`）。
> 旧 systemd 服务已 disable + stopped。

### 局域网连接（读写分离）

| 用途 | 地址 | 端口 |
|------|------|------|
| 写连接 | 10.167.2.175 | 26257 |
| 读连接 | 10.167.2.175 | 26267 |
| NodePort 写 | 10.167.2.175 | 30257 |
| NodePort 读 | 10.167.2.175 | 30267 |
| 管理界面 | http://10.167.2.175 | 30259 |
| K8s 内写 | cockroachdb.infra.svc | 26257 |
| K8s 内读 | cockroachdb-read.infra.svc | 26267 |
"""

PPT_CONTENT = [
    "JupyterHub v2.1 多语言审查+CRDB 确认",
    "",
    "【多语言代码审查】",
    "  Python: pycodestyle + pyflakes",
    "  Java: Checkstyle / javac",
    "  Go: gofmt + go vet",
    "  Rust: cargo clippy + rustfmt",
    "  C/C++: cppcheck",
    "  JavaScript: node --check",
    "  使用: python3 code_review.py <文件或目录>",
    "",
    "【指南分发规则】",
    "  教师/管理员: 教师版+学生版双指南",
    "  学生: 仅学生版指南",
    "  所有用户: 评分系统 + AI助手",
    "",
    "【CockroachDB K8s 部署确认】",
    "  3节点 StatefulSet (infra命名空间)",
    "  hostNetwork: true (非裸金属)",
    "  旧systemd已 disabled + stopped",
    "  Rancher可直接管理",
    "  读写分离: 26257(写) / 26267(读)",
    "",
    "【集群状态】",
    "  17个教师/管理员 | 20个分组",
    "  56个Notebook | 28个代码框架",
    "  3门课程 | 48核 CPU",
]

def update_ppts():
    """Update PPT files with new content."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    updated, skipped, errors = 0, 0, 0
    ppt_files = [f for f in glob.glob(os.path.join(BASE, "**", "*.pptx"), recursive=True)
                 if "_docx_work" not in f and "venv" not in f]

    for ppt_path in ppt_files:
        try:
            prs = Presentation(ppt_path)
            # Check if already has this slide
            has = False
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and "多语言审查" in shape.text_frame.text:
                        has = True
                        break
                if has: break
            if has:
                skipped += 1
                continue

            # Add slide
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
            tb.text_frame.text = PPT_CONTENT[0]
            p = tb.text_frame.paragraphs[0]
            p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F,0x49,0x7D)

            cb = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(6))
            ctf = cb.text_frame; ctf.word_wrap = True
            for i, line in enumerate(PPT_CONTENT[1:]):
                p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
                p.text = line
                if line.startswith("【"):
                    p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F,0x49,0x7D)
                else:
                    p.font.size = Pt(10); p.font.color.rgb = RGBColor(0x33,0x33,0x33)
            prs.save(ppt_path)
            updated += 1
        except Exception as e:
            errors += 1
    return updated, skipped, errors, len(ppt_files)

def update_mds():
    """Update MD files with new content."""
    updated, skipped = 0, 0
    md_files = []
    for root, dirs, files in os.walk(BASE):
        if "_docx_work" in root or "venv" in root or "node_modules" in root: continue
        for f in files:
            if f.endswith(".md"):
                md_files.append(os.path.join(root, f))

    for md_path in md_files:
        try:
            with open(md_path, "r", encoding="utf-8") as f:
                content = f.read()
            if "多语言代码审查系统" in content:
                skipped += 1
                continue
            with open(md_path, "a", encoding="utf-8") as f:
                f.write(GUIDE_APPEND)
            updated += 1
        except:
            pass
    return updated, skipped, len(md_files)

def update_dify_guides():
    """Update guides in D:\\dify-install."""
    updated = 0
    for fname in ["JUPYTERHUB-OPERATION-GUIDE.md", "JUPYTERHUB-STUDENT-GUIDE.md"]:
        path = os.path.join(DIFY, fname)
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
            if "多语言代码审查系统" not in content:
                with open(path, "a", encoding="utf-8") as f:
                    f.write(GUIDE_APPEND)
                updated += 1
    return updated

if __name__ == "__main__":
    print("=" * 70)
    print("Full Documentation Sync - Multi-language Review + CRDB K8s")
    print("=" * 70)

    # 1. Update PPTs
    print("\n--- PPT files ---")
    p, s, e, t = update_ppts()
    print(f"PPT: {p} updated, {s} skipped, {e} errors, {t} total")

    # 2. Update MDs
    print("\n--- Markdown files ---")
    mu, ms, mt = update_mds()
    print(f"MD: {mu} updated, {ms} skipped, {mt} total")

    # 3. Update Dify guides
    print("\n--- Dify guides ---")
    du = update_dify_guides()
    print(f"Dify guides: {du} updated")

    print(f"\n{'='*70}")
    print(f"SUMMARY: {p} PPT + {mu} MD + {du} guides = {p+mu+du} files updated")
    print(f"{'='*70}")
